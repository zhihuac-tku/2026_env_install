 # py code beginning

"""Document reader for CDC reporting materials.

Responsibilities:
- Read supported CDC reporting documents.
- Preserve useful document structure.
- Return a neutral ExtractedReportDocument.

This module does NOT:
- perform LLM analysis,
- determine CDC reporting requirements,
- perform FHIR mapping,
- perform IG gap analysis.
"""

from __future__ import annotations

import csv
import io
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field


# =========================================
# EXTRACTED DOCUMENT MODELS
# =========================================


class ExtractedTable(BaseModel):
    """One table extracted from a document."""

    name: str | None = None

    rows: list[list[str]] = Field(
        default_factory=list
    )


class ExtractedSheet(BaseModel):
    """One worksheet extracted from Excel."""

    name: str

    rows: list[list[str]] = Field(
        default_factory=list
    )


class ExtractedReportDocument(BaseModel):
    """Neutral representation of an uploaded report."""

    file_name: str

    file_type: str

    text: str = ""

    paragraphs: list[str] = Field(
        default_factory=list
    )

    tables: list[ExtractedTable] = Field(
        default_factory=list
    )

    sheets: list[ExtractedSheet] = Field(
        default_factory=list
    )

    extraction_notes: list[str] = Field(
        default_factory=list
    )


# =========================================
# HELPERS
# =========================================


def _clean_text(
    value: Any,
) -> str:
    """Convert a cell/value to clean text."""

    if value is None:
        return ""

    return str(value).strip()


def _rows_to_text(
    rows: list[list[str]],
) -> str:
    """Create readable text without losing rows."""

    lines: list[str] = []

    for row in rows:

        values = [
            value
            for value in row
            if value
        ]

        if values:
            lines.append(
                " | ".join(values)
            )

    return "\n".join(lines)


# =========================================
# ODT
# =========================================


def _read_odt(
    *,
    file_bytes: bytes,
    file_name: str,
) -> ExtractedReportDocument:

    try:
        from odf.opendocument import load
        from odf.table import (
            Table,
            TableCell,
            TableRow,
        )
        from odf.text import P
        from odf import teletype

    except ImportError as exc:
        raise RuntimeError(
            "ODT support requires odfpy."
        ) from exc

    document = load(
        io.BytesIO(file_bytes)
    )

    paragraphs: list[str] = []

    for paragraph in document.getElementsByType(
        P
    ):
        text = _clean_text(
            teletype.extractText(
                paragraph
            )
        )

        if text:
            paragraphs.append(text)

    tables: list[ExtractedTable] = []

    for table_index, table in enumerate(
        document.getElementsByType(Table),
        start=1,
    ):

        rows: list[list[str]] = []

        for row in table.getElementsByType(
            TableRow
        ):

            row_values: list[str] = []

            for cell in row.getElementsByType(
                TableCell
            ):

                cell_text = _clean_text(
                    teletype.extractText(
                        cell
                    )
                )

                row_values.append(
                    cell_text
                )

            if any(row_values):
                rows.append(row_values)

        if rows:
            tables.append(
                ExtractedTable(
                    name=f"Table {table_index}",
                    rows=rows,
                )
            )

    text_parts = list(paragraphs)

    for table in tables:
        table_text = _rows_to_text(
            table.rows
        )

        if table_text:
            text_parts.append(table_text)

    return ExtractedReportDocument(
        file_name=file_name,
        file_type="odt",
        text="\n".join(text_parts),
        paragraphs=paragraphs,
        tables=tables,
    )


# =========================================
# XLSX
# =========================================


def _read_xlsx(
    *,
    file_bytes: bytes,
    file_name: str,
) -> ExtractedReportDocument:

    try:
        from openpyxl import load_workbook

    except ImportError as exc:
        raise RuntimeError(
            "XLSX support requires openpyxl."
        ) from exc

    workbook = load_workbook(
        io.BytesIO(file_bytes),
        data_only=True,
        read_only=True,
    )

    sheets: list[ExtractedSheet] = []
    text_parts: list[str] = []

    for worksheet in workbook.worksheets:

        rows: list[list[str]] = []

        for row in worksheet.iter_rows(
            values_only=True
        ):

            values = [
                _clean_text(value)
                for value in row
            ]

            if any(values):
                rows.append(values)

        sheets.append(
            ExtractedSheet(
                name=worksheet.title,
                rows=rows,
            )
        )

        sheet_text = _rows_to_text(
            rows
        )

        if sheet_text:
            text_parts.append(
                f"[Sheet: {worksheet.title}]\n"
                f"{sheet_text}"
            )

    return ExtractedReportDocument(
        file_name=file_name,
        file_type="xlsx",
        text="\n\n".join(text_parts),
        sheets=sheets,
    )


# =========================================
# CSV
# =========================================


def _read_csv(
    *,
    file_bytes: bytes,
    file_name: str,
) -> ExtractedReportDocument:

    decoded_text: str | None = None

    for encoding in (
        "utf-8-sig",
        "utf-8",
        "cp950",
        "big5",
    ):

        try:
            decoded_text = (
                file_bytes.decode(
                    encoding
                )
            )
            break

        except UnicodeDecodeError:
            continue

    if decoded_text is None:
        raise ValueError(
            "Unable to decode CSV file."
        )

    reader = csv.reader(
        io.StringIO(decoded_text)
    )

    rows: list[list[str]] = []

    for row in reader:

        values = [
            _clean_text(value)
            for value in row
        ]

        if any(values):
            rows.append(values)

    return ExtractedReportDocument(
        file_name=file_name,
        file_type="csv",
        text=_rows_to_text(rows),
        tables=[
            ExtractedTable(
                name="CSV",
                rows=rows,
            )
        ],
    )


# =========================================
# XML
# =========================================


def _read_xml(
    *,
    file_bytes: bytes,
    file_name: str,
) -> ExtractedReportDocument:

    root = ET.fromstring(
        file_bytes
    )

    lines: list[str] = []

    for element in root.iter():

        value = _clean_text(
            element.text
        )

        if value:
            lines.append(
                f"{element.tag}: {value}"
            )

    return ExtractedReportDocument(
        file_name=file_name,
        file_type="xml",
        text="\n".join(lines),
        paragraphs=lines,
    )


# =========================================
# PDF
# =========================================


def _read_pdf(
    *,
    file_bytes: bytes,
    file_name: str,
) -> ExtractedReportDocument:

    try:
        from pypdf import PdfReader

    except ImportError as exc:
        raise RuntimeError(
            "PDF support requires pypdf."
        ) from exc

    reader = PdfReader(
        io.BytesIO(file_bytes)
    )

    pages: list[str] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):

        page_text = _clean_text(
            page.extract_text()
        )

        if page_text:
            pages.append(
                f"[Page {page_number}]\n"
                f"{page_text}"
            )

    extraction_notes: list[str] = []

    if not pages:
        extraction_notes.append(
            "No embedded PDF text was found. "
            "The PDF may be image-based and may "
            "require OCR."
        )

    return ExtractedReportDocument(
        file_name=file_name,
        file_type="pdf",
        text="\n\n".join(pages),
        paragraphs=pages,
        extraction_notes=extraction_notes,
    )


# =========================================
# PUBLIC READER
# =========================================


def read_cdc_report(
    *,
    file_bytes: bytes,
    file_name: str,
    file_type: str | None = None,
) -> ExtractedReportDocument:
    """Read a supported CDC reporting document."""

    if not file_bytes:
        raise ValueError(
            "CDC report file is empty."
        )

    normalized_type = (
        file_type
        or Path(file_name).suffix.lstrip(".")
    ).lower().strip()

    readers = {
        "odt": _read_odt,
        "xlsx": _read_xlsx,
        "csv": _read_csv,
        "xml": _read_xml,
        "pdf": _read_pdf,
    }

    reader = readers.get(
        normalized_type
    )

    if reader is None:
        raise ValueError(
            "Unsupported CDC report file type: "
            f"{normalized_type}"
        )

    return reader(
        file_bytes=file_bytes,
        file_name=file_name,
    )

